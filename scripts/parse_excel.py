#!/usr/bin/env python3
import sys
import openpyxl
import json
import os

if len(sys.argv) < 2:
    print(json.dumps({'success': False, 'message': 'File path required'}))
    sys.exit(1)

excel_path = sys.argv[1]
output_path = sys.argv[2] if len(sys.argv) > 2 else '/home/www/cammon/data/cctv_config.json'

if not os.path.exists(excel_path):
    print(json.dumps({'success': False, 'message': f'File not found: {excel_path}'}))
    sys.exit(1)

try:
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb.active

    existing_map = {}
    if os.path.exists(output_path):
        try:
            prev_list = json.load(open(output_path, 'r', encoding='utf-8'))
            for c in prev_list:
                existing_map[c['id']] = c
        except Exception:
            pass

    # Find headers
    header_map = {}
    start_row = 1
    for r in range(1, 10):
        row_vals = [sheet.cell(row=r, column=c).value for c in range(1, 20)]
        if "순번" in str(row_vals):
            for idx, val in enumerate(row_vals):
                if val:
                    header_map[str(val).strip()] = idx + 1
            start_row = r + 1
            break
            
    if not header_map:
        # Fallback to column index if headers not found
        header_map = {"순번": 1, "경기장 번호": 2, "경기장명": 3, "경기직종명": 4, "DDNS명": 5}
        start_row = 2

    cams = []
    
    for r in range(start_row, sheet.max_row + 1):
        cam_num = sheet.cell(row=r, column=header_map.get("순번", 1)).value
        stadium_code = sheet.cell(row=r, column=header_map.get("경기장 번호", 2)).value
        stadium_name = sheet.cell(row=r, column=header_map.get("경기장명", 3)).value
        category = sheet.cell(row=r, column=header_map.get("경기직종명", 4)).value
        ddns = sheet.cell(row=r, column=header_map.get("DDNS명", 5)).value

        if cam_num is None or str(cam_num).strip() == '':
            continue
        try:
            cam_num = int(cam_num)
        except ValueError:
            continue

        s_code = str(stadium_code).strip() if stadium_code else ''
        s_name = str(stadium_name).strip() if stadium_name else ''
        stadium_full = f'{s_name} ({s_code})' if s_name and s_code else (s_name or s_code)

        cam_id = f'cam{cam_num:02d}'
        prev_data = existing_map.get(cam_id, {})

        # 엑셀이 권위: 경기장/종목/DDNS. IP는 cctv_ips.json(배포정보)이 권위 — 여기서 건드리지 않음.
        entry = {
            'id': cam_id,
            'num': cam_num,
            'stadium': stadium_full,
            'category': str(category).strip() if category else '',
            'ddns': str(ddns).strip() if ddns else '',
            'status': prev_data.get('status', 'STOPPED'),
            'alert_enabled': prev_data.get('alert_enabled', True),
            'alert_start': prev_data.get('alert_start', '08:00'),
            'alert_end': prev_data.get('alert_end', '18:00'),
        }
        # 운영 메타는 유지 (엑셀에 없는 필드)
        if prev_data.get('host_name'):
            entry['host_name'] = prev_data['host_name']
        if 'suppress_until' in prev_data:
            entry['suppress_until'] = prev_data.get('suppress_until')

        cams.append(entry)

    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(cams, f, ensure_ascii=False, indent=2)

    removed = len(existing_map) - len({c['id'] for c in cams} & set(existing_map.keys()))
    msg = f'Successfully updated {len(cams)} cameras from Excel file'
    if removed > 0:
        msg += f' ({removed} removed from monitoring scope)'

    print(json.dumps({
        'success': True,
        'message': msg,
        'count': len(cams),
        'removed': removed,
    }))

except Exception as e:
    print(json.dumps({'success': False, 'message': f'Error parsing Excel: {str(e)}'}))
